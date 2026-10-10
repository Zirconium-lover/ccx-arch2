#!/usr/bin/env python3
"""Three interchangeable decks from the same manim-slides scenes.

    make_decks.py OUT_STEM Scene1 [Scene2 ...]      (run in the anim/ folder)

  OUT_STEM_video.pptx  one slide per animation step, the step's MP4 starts by
                       itself when the slide is shown (a click = next step).
                       Poster = first frame, so nothing blinks; one static
                       slide with the final frame closes every scene, so if a
                       PC cannot play video the deck degrades to a flip-book.
  OUT_STEM_gif.pptx    the same slides with GIFs instead of MP4: no codec is
                       involved, plays in any PowerPoint / Impress / Keynote.
  OUT_STEM.pdf         the final frame of every step: the last-resort copy.
"""
import json, subprocess, sys, tempfile
from pathlib import Path

import pptx
from pptx.util import Emu
from PIL import Image
from manim_slides.convert import _auto_play_pptx_media

FPS_GIF = 25


def frame(video, which, out):
    args = ['ffmpeg', '-v', 'error', '-y']
    args += (['-sseof', '-0.05'] if which == 'last' else [])
    subprocess.run(args + ['-i', str(video), '-frames:v', '1', str(out)], check=True)


def gif(video, out, width):
    vf = (f'fps={FPS_GIF},scale={width}:-1:flags=lanczos,split[a][b];'
          '[a]palettegen=max_colors=128:stats_mode=full[p];'
          '[b][p]paletteuse=dither=none')
    # -loop -1: play once and stay on the last frame
    subprocess.run(['nice', 'ffmpeg', '-v', 'error', '-y', '-i', str(video), '-vf', vf,
                    '-loop', '-1', str(out)], check=True)


def main():
    stem, scenes = sys.argv[1], sys.argv[2:]
    steps = []                       # (video, is_last_of_scene)
    for sc in scenes:
        cfg = json.load(open(f'slides/{sc}.json'))
        W, H = cfg['resolution']
        files = [Path(s['file']) for s in cfg['slides']]
        steps += [(f, i == len(files) - 1) for i, f in enumerate(files)]

    tmp = Path(tempfile.mkdtemp())
    decks = {'video': pptx.Presentation(), 'gif': pptx.Presentation()}
    for p in decks.values():
        p.slide_width, p.slide_height = Emu(W * 9525), Emu(H * 9525)
    pdf_pages = []
    sw, sh = Emu(W * 9525), Emu(H * 9525)
    for n, (v, last) in enumerate(steps):
        first_png, last_png, g = tmp / f'{n}_a.png', tmp / f'{n}_z.png', tmp / f'{n}.gif'
        frame(v, 'first', first_png)
        frame(v, 'last', last_png)
        gif(v, g, W)
        pdf_pages.append(Image.open(last_png).convert('RGB'))

        s = decks['video'].slides.add_slide(decks['video'].slide_layouts[6])
        m = s.shapes.add_movie(str(v), 0, 0, sw, sh, poster_frame_image=str(first_png),
                               mime_type='video/mp4')
        _auto_play_pptx_media(m)
        s = decks['gif'].slides.add_slide(decks['gif'].slide_layouts[6])
        s.shapes.add_picture(str(g), 0, 0, sw, sh)
        if last:
            for p in decks.values():
                p.slides.add_slide(p.slide_layouts[6]).shapes.add_picture(
                    str(last_png), 0, 0, sw, sh)
    for k, p in decks.items():
        p.save(f'{stem}_{k}.pptx')
    pdf_pages[0].save(f'{stem}.pdf', save_all=True, append_images=pdf_pages[1:])
    for f in (f'{stem}_video.pptx', f'{stem}_gif.pptx', f'{stem}.pdf'):
        print('%-28s %6.1f MB' % (f, Path(f).stat().st_size / 1e6))


if __name__ == '__main__':
    main()
